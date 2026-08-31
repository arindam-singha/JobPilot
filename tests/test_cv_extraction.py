from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter

from app.cv.extraction import DOCXTextExtractor, PDFTextExtractor, extract_document_text


@pytest.fixture
def simple_pdf_bytes() -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    page.merge_transformed_page(page, None)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


@pytest.fixture
def simple_docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("First line")
    document.add_paragraph("Second line")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_pdf_extractor_returns_text(simple_pdf_bytes):
    result = PDFTextExtractor().extract(simple_pdf_bytes)

    assert isinstance(result.text, str)
    assert result.page_count == 1 or result.page_count is None


def test_docx_extractor_returns_text(simple_docx_bytes):
    result = DOCXTextExtractor().extract(simple_docx_bytes)

    assert result.text == "First line\nSecond line"
    assert result.paragraph_count == 2


def test_extract_document_text_dispatches_pdf(simple_pdf_bytes):
    result = extract_document_text(simple_pdf_bytes, "application/pdf")

    assert isinstance(result.text, str)


def test_extract_document_text_dispatches_docx(simple_docx_bytes):
    result = extract_document_text(simple_docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    assert result.text == "First line\nSecond line"


def test_extract_document_text_raises_for_unsupported_mime_type():
    with pytest.raises(ValueError):
        extract_document_text(b"random bytes", "text/plain")
